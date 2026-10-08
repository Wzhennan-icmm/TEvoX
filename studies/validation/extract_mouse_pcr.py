#!/usr/bin/env python3
"""Join Table S2/S3 from PMID 37228752 and retrieve mm39 reference alleles.

Requires openpyxl. The article supplementary ZIP is publicly available at
https://www.ebi.ac.uk/europepmc/webservices/rest/PMC10203049/supplementaryFiles
Table S2 is mmc3.xlsx and Table S3 is mmc4.xlsx. No prediction file is read.
"""
import argparse
import json
from pathlib import Path
import urllib.request

import openpyxl
from mouse_pcr_truth import sha256


def main(args):
    args.output.mkdir(parents=True, exist_ok=True)
    pcr = {}
    replicates = {}
    pcr_rows = 0
    book = openpyxl.load_workbook(args.table_s3, read_only=True, data_only=True)
    for sheet in book:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        if sheet.title == 'PCR_validations':
            header = next(rows)
        if 'CAST' not in header or header[0] != 'ID':
            raise ValueError('unexpected PCR table schema')
        for row in rows:
            if not row[0]:
                continue
            pcr_rows += 1
            record = dict(dict(zip(header, row)), sheet=sheet.title)
            replicates.setdefault(row[0], []).append(record)
            if row[0] in pcr and pcr[row[0]]['CAST'] != record['CAST']:
                raise ValueError('conflicting replicate CAST PCR labels')
            pcr.setdefault(row[0], record)
    book.close()
    book = openpyxl.load_workbook(args.table_s2, read_only=True, data_only=True)
    rows = book['INS-DEL'].iter_rows(values_only=True)
    header = next(rows)
    matched = []
    for row in rows:
        if row[0] in pcr:
            matched.append({'pcr': pcr[row[0]], 'sv': dict(zip(header, row))})
    book.close()
    ids = [r['sv']['ID'] for r in matched]
    if len(ids) != len(set(ids)):
        raise ValueError('duplicate SV variant ID')
    candidates = []
    for record in matched:
        sv = record['sv']
        if sv['SVTYPE'] != 'DEL' or not sv['TE_TYPE'] or record['pcr']['CAST'] not in (0, 1):
            continue
        if sv['END'] - sv['POS'] != sv['SVLEN'] or sv['POS'] < 50:
            raise ValueError('unexpected deletion coordinates')
        start, end = sv['POS'] - 50, sv['END'] + 50
        url = 'https://api.genome.ucsc.edu/getData/sequence?genome=mm39;chrom={};start={};end={}'.format(sv['CHROM'], start, end)
        path = args.output / (sv['ID'] + '.json')
        if not path.exists():
            with urllib.request.urlopen(url, timeout=60) as response:
                data = response.read()
            parsed = json.loads(data)
            if (parsed['genome'], parsed['chrom'], parsed['start'], parsed['end'], len(parsed['dna'])) != ('mm39', sv['CHROM'], start, end, end - start):
                raise ValueError('reference response mismatch')
            path.write_bytes(data)
        candidates.append(dict(record, reference={'url': url, 'path': path.name,
                                                  'sha256': sha256(path), 'start': start, 'end': end, 'flank': 50}))
    output = {'publication': 'PMID:37228752; DOI:10.1016/j.xgen.2023.100291',
              'input_sha256': {'mmc3.xlsx': sha256(args.table_s2), 'mmc4.xlsx': sha256(args.table_s3)},
              'pcr_rows': pcr_rows, 'pcr_unique_variants': len(pcr), 'joined_records': len(matched),
              'repeated_variant_rows': {key: rows for key, rows in replicates.items() if len(rows) > 1},
              'unmatched_pcr_ids': sorted(set(pcr) - set(ids)),
              'protocol': 'PCR DEL with TE_TYPE; CAST observed 0/1; 50 bp flanks; exact unique full REF+flank transfer; >=.95 TE coverage and >=.8 DEL coverage; selected without inspecting predictions',
              'candidates': candidates}
    (args.output / 'candidates.json').write_text(json.dumps(output, indent=2) + '\n')
    print('PCR records:', len(pcr), 'joined:', len(matched), 'eligible DEL:', len(candidates))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--table-s2', type=Path, required=True)
    parser.add_argument('--table-s3', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    main(parser.parse_args())
